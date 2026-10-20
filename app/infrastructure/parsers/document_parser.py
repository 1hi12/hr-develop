"""
DocumentParser — implements IDocumentParser.
Supports: PDF (pdfplumber + pymupdf fallback), DOCX, Image (OCR), TXT.
"""

import io
import logging
import re
from pathlib import Path

from app.core.ports.document_parser import IDocumentParser

logger = logging.getLogger(__name__)

ALLOWED_TYPES = {
    "pdf":  "pdf",
    "docx": "docx",
    "doc":  "docx",
    "png":  "image",
    "jpg":  "image",
    "jpeg": "image",
    "webp": "image",
    "tiff": "image",
    "txt":  "text",
}


class DocumentParser(IDocumentParser):

    def parse(self, file_bytes: bytes, filename: str) -> str:
        suffix    = Path(filename).suffix.lower().lstrip(".")
        file_type = ALLOWED_TYPES.get(suffix)

        if not file_type:
            raise ValueError(
                f"Unsupported format: .{suffix}. "
                f"Accepted: {list(ALLOWED_TYPES.keys())}"
            )
        if len(file_bytes) < 100:
            raise ValueError("Empty or corrupt file")

        size_mb = len(file_bytes) / (1024 * 1024)
        if size_mb > 20:
            raise ValueError(f"File too large: {size_mb:.1f} MB (max 20 MB)")

        logger.info(
            "[Parser] Loading: %s | type=%s | %.2f MB", filename, file_type, size_mb
        )

        extractors = {
            "pdf":   self._extract_pdf,
            "docx":  self._extract_docx,
            "image": self._extract_image,
            "text":  self._extract_plain_text,
        }

        raw_text = extractors[file_type](file_bytes)
        raw_text = self._clean_text(raw_text)

        if len(raw_text.strip()) < 50:
            logger.warning(
                "[Parser] Very short text: %d chars for %s", len(raw_text), filename
            )

        logger.info("[Parser] Extracted %d characters", len(raw_text))
        return raw_text

    # ── Extractors ────────────────────────────────────────────────────────────

    def _extract_pdf(self, file_bytes: bytes) -> str:
        try:
            import pdfplumber
            pages_text = []
            with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                for i, page in enumerate(pdf.pages):
                    text = page.extract_text(x_tolerance=2, y_tolerance=2)
                    if text and text.strip():
                        pages_text.append(f"[Page {i + 1}]\n{text}")
            result = "\n\n".join(pages_text)
            if result.strip():
                logger.info("[Parser] pdfplumber OK — %d chars", len(result))
                return result
        except Exception as e:
            logger.warning("[Parser] pdfplumber failed: %s", e)

        try:
            import fitz
            doc = fitz.open(stream=file_bytes, filetype="pdf")
            pages_text = []
            for i, page in enumerate(doc):
                text = page.get_text("text")
                if text.strip():
                    pages_text.append(f"[Page {i + 1}]\n{text}")
            result = "\n\n".join(pages_text)
            logger.info("[Parser] pymupdf OK — %d chars", len(result))
            return result
        except Exception as e:
            logger.error("[Parser] pymupdf failed: %s", e)
            raise RuntimeError(f"Cannot extract text from PDF: {e}")

    def _extract_pdf_with_formatting(self, file_bytes: bytes) -> str:
        """
        Extract PDF text with Markdown formatting (bold, lists, etc.)
        Uses pymupdf to analyze text styling and structure.
        """
        try:
            import fitz
            doc = fitz.open(stream=file_bytes, filetype="pdf")
            markdown_lines = []
            
            for page_num, page in enumerate(doc):
                markdown_lines.append(f"[Page {page_num + 1}]")
                
                # Get text blocks with formatting information
                blocks = page.get_text("dict")["blocks"]
                
                for block in blocks:
                    if block.get("type") != 0:  # Skip non-text blocks
                        continue
                    
                    for line in block.get("lines", []):
                        line_text = ""
                        
                        for span in line.get("spans", []):
                            text = span.get("text", "").strip()
                            if not text:
                                continue
                            
                            flags = span.get("flags", 0)
                            font_size = span.get("size", 11)
                            
                            # Check formatting flags
                            is_bold = bool(flags & 2**4)  # Bold flag
                            is_italic = bool(flags & 2**1)  # Italic flag
                            
                            # Detect headings by font size
                            if font_size > 14:
                                text = f"## {text}"
                            elif font_size > 12:
                                text = f"### {text}"
                            elif is_bold and is_italic:
                                text = f"***{text}***"
                            elif is_bold:
                                text = f"**{text}**"
                            elif is_italic:
                                text = f"*{text}*"
                            
                            line_text += text + " "
                        
                        line_text = line_text.strip()
                        
                        # Detect bullet points and numbered lists
                        if line_text:
                            # Numbered lists: "1.", "2.", etc.
                            if re.match(r'^\d+\.\s+', line_text):
                                pass  # Keep as is
                            # Bullet points: •, -, *, etc.
                            elif re.match(r'^[•\-\*○▪]\s+', line_text):
                                line_text = re.sub(r'^[•○▪]\s+', '- ', line_text)
                            
                            markdown_lines.append(line_text)
                
                markdown_lines.append("")  # Empty line between pages
            
            result = "\n".join(markdown_lines)
            logger.info("[Parser] PDF with formatting extracted — %d chars", len(result))
            return result
            
        except Exception as e:
            logger.warning("[Parser] Formatted extraction failed: %s, falling back to plain text", e)
            # Fallback to plain text extraction
            return self._extract_pdf(file_bytes)

    def _extract_docx(self, file_bytes: bytes) -> str:
        try:
            from docx import Document
            doc   = Document(io.BytesIO(file_bytes))
            parts = []
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    parts.append(text)
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(
                        cell.text.strip() for cell in row.cells if cell.text.strip()
                    )
                    if row_text:
                        parts.append(row_text)
            result = "\n".join(parts)
            logger.info("[Parser] DOCX OK — %d chars", len(result))
            return result
        except Exception as e:
            logger.error("[Parser] DOCX error: %s", e)
            raise RuntimeError(f"Cannot extract DOCX text: {e}")

    def _extract_image(self, file_bytes: bytes) -> str:
        try:
            import pytesseract
            from PIL import Image, ImageFilter, ImageEnhance
            image    = Image.open(io.BytesIO(file_bytes)).convert("L")
            image    = image.filter(ImageFilter.SHARPEN)
            image    = ImageEnhance.Contrast(image).enhance(2.0)
            text     = pytesseract.image_to_string(
                image, lang="fra+eng", config="--psm 6 --oem 3"
            )
            logger.info("[Parser] OCR OK — %d chars", len(text))
            return text
        except Exception as e:
            logger.error("[Parser] OCR error: %s", e)
            raise RuntimeError(f"Cannot extract text from image: {e}")

    def _extract_plain_text(self, file_bytes: bytes) -> str:
        for encoding in ("utf-8", "latin-1", "cp1252"):
            try:
                return file_bytes.decode(encoding)
            except UnicodeDecodeError:
                continue
        return file_bytes.decode("utf-8", errors="replace")

    def _clean_text(self, text: str) -> str:
        text = re.sub(r"[^\x20-\x7E\xA0-\xFF\n\t]", " ", text)
        text = text.replace("\t", " ")
        text = re.sub(r" {3,}", " ", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        return text.strip()
